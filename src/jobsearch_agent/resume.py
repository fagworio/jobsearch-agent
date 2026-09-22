"""Geração, validação factual/ATS e renderização de currículo."""

from __future__ import annotations

import html
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from .llm import LLMError, LLMProvider, LLMRequest
from .models import CareerProfile, Fact, Job, Resume, ResumeClaim, ResumeStrategy, ValidationResult
from .skills import SkillRegistry

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError:  # pragma: no cover - fallback do ambiente mínimo
    Document = None
    Inches = Pt = None


def _statement(fact: Fact, language: str) -> str:
    return fact.statements.get(language) or fact.statements.get("en-US") or next(iter(fact.statements.values()))


STOPWORDS = {
    "a", "an", "and", "as", "at", "by", "com", "da", "de", "do", "e", "for", "in", "na", "no", "of", "on", "or", "para", "the", "to", "um", "uma", "with",
}
NUMBER_RE = re.compile(r"(?<![\w])(?:\$|€|R\$)?\s*\d+(?:[.,]\d+)?\s*%?(?![\w])")


def _tokens(text: str) -> list[str]:
    return re.findall(r"[a-záéíóúãõç0-9][a-záéíóúãõç0-9+.#-]*", text.lower())


def _token_supported(token: str, corpus_tokens: set[str]) -> bool:
    if token in STOPWORDS or len(token) <= 2:
        return True
    if token in corpus_tokens:
        return True
    # Permite flexões simples sem liberar novos conceitos: developer/developed,
    # integrations/integration e equivalentes compartilham raiz lexical.
    if len(token) >= 7 and any(len(other) >= 7 and token[:6] == other[:6] for other in corpus_tokens):
        return True
    return False


def rank_fact_ids(job: Job, strategy: ResumeStrategy, profile: CareerProfile, facts: dict[str, Fact]) -> list[str]:
    registry = SkillRegistry.load()
    wanted = strategy.keywords + strategy.focus + strategy.secondary
    target = f"{job.title} {job.description}".lower()
    scored: list[tuple[float, int, str]] = []
    order = 0
    for experience in profile.experiences:
        for fact_id in experience.fact_ids:
            fact = facts.get(fact_id)
            if not fact:
                continue
            score = 0.0
            for tag in fact.tags:
                if any(registry.matches(tag, term) for term in wanted):
                    score += 4.0
                if tag.lower() in target:
                    score += 1.0
            statement = _statement(fact, strategy.language).lower()
            score += sum(0.5 for term in wanted if term.lower() in statement or term.lower() in target and term.lower() in statement)
            scored.append((score, order, fact_id))
            order += 1
    return [fact_id for _, _, fact_id in sorted(scored, key=lambda item: (-item[0], item[1]))]


def select_fact_ids(job: Job, strategy: ResumeStrategy, profile: CareerProfile, facts: dict[str, Fact], max_facts: int = 8) -> list[str]:
    ranked = rank_fact_ids(job, strategy, profile, facts)
    selected = ranked[:max_facts]
    if selected:
        return selected
    return [fact_id for experience in profile.experiences for fact_id in experience.fact_ids if fact_id in facts][:max_facts]


def cluster_fact_ids(fact_ids: list[str], facts: dict[str, Fact], max_per_claim: int = 2) -> list[list[str]]:
    """Agrupa facts pequenos sem jamais descartar um ID selecionado."""
    clusters: list[list[str]] = []
    for fact_id in fact_ids:
        tags = {tag.lower() for tag in facts[fact_id].tags}
        target = next((cluster for cluster in clusters if len(cluster) < max_per_claim and tags & {tag.lower() for item in cluster for tag in facts[item].tags}), None)
        if target is None:
            clusters.append([fact_id])
        else:
            target.append(fact_id)
    return clusters


def rewrite_claim(job: Job, strategy: ResumeStrategy, fact_ids: list[str], facts: dict[str, Fact], provider: LLMProvider | None = None, fallback_events: list[dict[str, str]] | None = None) -> tuple[str, list[str]]:
    """Monta um claim contextual preservando todos os facts do grupo.

    Quando configurado, o LLM recebe somente os facts selecionados e deve
    devolver ``text`` + ``supported_by``. Qualquer resposta que omita um fact
    selecionado ou não tenha estrutura válida cai no compositor determinístico.
    """
    language = strategy.language
    statements = [_statement(facts[fact_id], language) for fact_id in fact_ids]
    if provider and fact_ids:
        try:
            response = provider.complete(LLMRequest(
                system="Rewrite only the supplied facts for the target role. Return JSON with text and supported_by. Never add metrics, technologies, entities or responsibilities.",
                user="\n".join([
                    f"Target role: {job.title}",
                    f"Language: {language}",
                    f"Focus: {', '.join(strategy.focus + strategy.secondary)}",
                    "Facts:",
                    *[f"{fact_id}: {_statement(facts[fact_id], language)}" for fact_id in fact_ids],
                ]),
                schema="{text: string, supported_by: string[]}",
            ))
            text = response.get("text") if isinstance(response, dict) else None
            support = response.get("supported_by") if isinstance(response, dict) else None
            if isinstance(text, str) and text.strip() and isinstance(support, list) and set(support) == set(fact_ids):
                return text.strip(), list(support)
            if fallback_events is not None:
                fallback_events.append({"reason": "invalid_response", "provider": type(provider).__name__})
        except LLMError:
            if fallback_events is not None:
                fallback_events.append({"reason": "llm_error", "provider": type(provider).__name__})
    if len(statements) > 1:
        return "; ".join(statement.rstrip(".") for statement in statements) + ".", fact_ids
    return statements[0] if statements else "", fact_ids


def generate_resume(job: Job, strategy: ResumeStrategy, profile: CareerProfile, facts: dict[str, Fact], selected_ids: list[str], provider: LLMProvider | None = None, fallback_events: list[dict[str, str]] | None = None) -> Resume:
    language = strategy.language
    claims: list[ResumeClaim] = []
    summary = profile.professional_summary.get(language) or profile.professional_summary.get("en-US", "")
    summary_ids = [fact_id for fact_id in profile.summary_fact_ids.get(language, profile.summary_fact_ids.get("en-US", [])) if fact_id in facts]
    if summary_ids and provider:
        rewritten_summary, support = rewrite_claim(job, strategy, summary_ids, facts, provider, fallback_events)
        if rewritten_summary:
            summary, summary_ids = rewritten_summary, support
    if summary:
        claims.append(ResumeClaim(summary, summary_ids, bool(summary_ids)))
    experience_rows: list[dict[str, Any]] = []
    for experience in profile.experiences:
        exp_facts = [fact_id for fact_id in experience.fact_ids if fact_id in selected_ids]
        if not exp_facts:
            continue
        grouped_ids = cluster_fact_ids(exp_facts, facts)
        bullets: list[str] = []
        for group in grouped_ids:
            bullet, support = rewrite_claim(job, strategy, group, facts, provider, fallback_events)
            bullets.append(bullet)
            claims.append(ResumeClaim(bullet, support, True))
        experience_rows.append({"company": experience.company, "role": experience.role, "start_date": experience.start_date, "end_date": experience.end_date, "bullets": bullets, "fact_ids": exp_facts})
    education_rows: list[dict[str, Any]] = []
    for item in profile.education:
        fact_ids = [fact_id for fact_id in item.fact_ids if fact_id in facts]
        if not fact_ids:
            continue
        credential = item.credential.get(language) or item.credential.get("en-US") or next(iter(item.credential.values()), "")
        field_of_study = item.field_of_study.get(language) or item.field_of_study.get("en-US") or next(iter(item.field_of_study.values()), "")
        education_rows.append({
            "institution": item.institution,
            "credential": credential,
            "field_of_study": field_of_study,
            "start_date": item.start_date,
            "end_date": item.end_date,
            "fact_ids": fact_ids,
        })
        for fact_id in fact_ids:
            claims.append(ResumeClaim(_statement(facts[fact_id], language), [fact_id], True))
    return Resume(
        id=f"resume-{job.id}-{language}", job_id=job.id, language=language,
        header={"name": profile.identity.get("name", ""), "email": profile.identity.get("email", ""), "location": profile.identity.get("location", "")},
        summary=summary, skills=strategy.keywords, experience=experience_rows, education=education_rows, claims=claims,
    )


def validate_facts(resume: Resume, facts: dict[str, Fact]) -> ValidationResult:
    errors: list[str] = []
    details: dict[str, Any] = {"unsupported_claim_atoms": [], "claims": []}
    registry = SkillRegistry.load()
    for claim in resume.claims:
        supported = [fact_id for fact_id in claim.supported_by if fact_id in facts and facts[fact_id].verified]
        if not supported:
            claim.valid = False
            errors.append(f"unsupported claim: {claim.claim}")
            details["unsupported_claim_atoms"].append({"claim": claim.claim, "atoms": ["missing_fact_support"]})
            continue
        corpus = " ".join(_statement(facts[fact_id], resume.language).lower() for fact_id in supported)
        corpus_tokens = set(_tokens(corpus))
        unsupported_atoms: list[str] = []
        claim_numbers = {number.replace(" ", "") for number in NUMBER_RE.findall(claim.claim)}
        corpus_numbers = {number.replace(" ", "") for number in NUMBER_RE.findall(corpus)}
        unsupported_atoms.extend(f"number:{number}" for number in sorted(claim_numbers - corpus_numbers))
        unsupported_atoms.extend(token for token in _tokens(claim.claim) if not _token_supported(token, corpus_tokens))
        claim_skills = set(registry.extract(claim.claim))
        corpus_skills = set(registry.extract(corpus))
        unsupported_atoms.extend(f"technology:{skill}" for skill in sorted(claim_skills - corpus_skills))
        if unsupported_atoms:
            claim.valid = False
            errors.append(f"claim does not match supporting facts: {claim.claim}")
            details["unsupported_claim_atoms"].append({"claim": claim.claim, "atoms": sorted(set(unsupported_atoms))})
        details["claims"].append({"claim": claim.claim, "supported_by": supported, "valid": claim.valid})
    return ValidationResult(not errors, "RESUME_VALIDATION_FAILED" if errors else "OK", errors, details=details)


def validate_ats(resume: Resume) -> ValidationResult:
    errors: list[str] = []
    if not resume.header.get("name"):
        errors.append("contact name is missing")
    if not resume.summary:
        errors.append("summary is missing")
    if len(resume.experience) == 0:
        errors.append("experience is missing")
    if any(not row.get("bullets") for row in resume.experience):
        errors.append("experience entry has no bullets")
    return ValidationResult(not errors, "ATS_VALIDATION_FAILED" if errors else "OK", errors, details={"template": resume.template, "one_column": True})


def render_text(resume: Resume) -> str:
    labels = {
        "pt-BR": {"summary": "Resumo profissional", "skills": "Competências", "experience": "Experiência profissional", "education": "Formação", "present": "Atual"},
        "en-US": {"summary": "Summary", "skills": "Skills", "experience": "Experience", "education": "Education", "present": "Present"},
    }.get(resume.language, {"summary": "Summary", "skills": "Skills", "experience": "Experience", "education": "Education", "present": "Present"})
    lines = [resume.header.get("name", ""), resume.header.get("email", ""), resume.header.get("location", ""), "", labels["summary"], resume.summary, "", labels["skills"], ", ".join(resume.skills), "", labels["experience"]]
    for row in resume.experience:
        dates = f"{row['start_date']} - {row['end_date'] or labels['present']}"
        lines.extend([f"{row['role']} | {row['company']} | {dates}"] + [f"- {bullet}" for bullet in row["bullets"]] + [""])
    if resume.education:
        lines.extend([labels["education"]])
        for row in resume.education:
            dates = " - ".join(value for value in (row.get("start_date", ""), row.get("end_date", "")) if value)
            credential = ", ".join(value for value in (row.get("credential", ""), row.get("field_of_study", "")) if value)
            lines.extend([f"{credential} | {row['institution']}" + (f" | {dates}" if dates else ""), ""])
    return "\n".join(line for line in lines if line is not None).strip() + "\n"


def _xml_text(text: str) -> str:
    return html.escape(text, quote=False).replace("\n", "</w:t></w:r><w:r><w:br/><w:t>")


def _render_docx_python_docx(resume: Resume, target: Path) -> Path:
    labels = {
        "pt-BR": {"summary": "Resumo profissional", "skills": "Competências", "experience": "Experiência profissional", "education": "Formação", "present": "Atual"},
        "en-US": {"summary": "Summary", "skills": "Skills", "experience": "Experience", "education": "Education", "present": "Present"},
    }.get(resume.language, {"summary": "Summary", "skills": "Skills", "experience": "Experience", "education": "Education", "present": "Present"})
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    document.add_heading(resume.header.get("name", ""), level=0)
    contact = " | ".join(value for value in (resume.header.get("email", ""), resume.header.get("location", "")) if value)
    if contact:
        document.add_paragraph(contact)
    document.add_heading(labels["summary"], level=1)
    document.add_paragraph(resume.summary)
    document.add_heading(labels["skills"], level=1)
    document.add_paragraph(", ".join(resume.skills))
    document.add_heading(labels["experience"], level=1)
    for row in resume.experience:
        document.add_heading(f"{row['role']} | {row['company']}", level=2)
        document.add_paragraph(f"{row['start_date']} - {row['end_date'] or labels['present']}")
        for bullet in row["bullets"]:
            document.add_paragraph(bullet, style="List Bullet")
    if resume.education:
        document.add_heading(labels["education"], level=1)
        for row in resume.education:
            credential = ", ".join(value for value in (row.get("credential", ""), row.get("field_of_study", "")) if value)
            document.add_heading(f"{credential} | {row['institution']}", level=2)
            dates = " - ".join(value for value in (row.get("start_date", ""), row.get("end_date", "")) if value)
            if dates:
                document.add_paragraph(dates)
    document.save(target)
    return target


def _render_docx_compat(resume: Resume, target: Path) -> Path:
    """Fallback mínimo para o ambiente sem python-docx; produção usa a biblioteca."""
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    paragraphs = []
    for line in render_text(resume).splitlines():
        paragraphs.append(f"<w:p><w:r><w:t xml:space=\"preserve\">{_xml_text(line)}</w:t></w:r></w:p>")
    document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{''.join(paragraphs)}<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1080" w:right="1080" w:bottom="1080" w:left="1080"/></w:sectPr></w:body></w:document>'''
    content_types = '''<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return target


def render_docx(resume: Resume, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    if Document is not None:
        return _render_docx_python_docx(resume, target)
    return _render_docx_compat(resume, target)


def render_pdf_from_docx(docx_path: str | Path, pdf_path: str | Path) -> Path:
    target = Path(pdf_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        raise RuntimeError("LibreOffice is required to render PDF")
    with tempfile.TemporaryDirectory(prefix="jobsearch-pdf-") as directory:
        result = subprocess.run([libreoffice, "--headless", "--convert-to", "pdf", "--outdir", directory, str(docx_path)], capture_output=True, text=True, timeout=60)
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "PDF conversion failed")
        generated = Path(directory) / (Path(docx_path).stem + ".pdf")
        if not generated.exists():
            raise RuntimeError("LibreOffice did not produce a PDF")
        target.write_bytes(generated.read_bytes())
    return target
