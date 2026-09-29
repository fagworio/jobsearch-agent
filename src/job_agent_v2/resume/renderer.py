"""Renderização ATS de uma coluna para DOCX e PDF."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from .models import ResumeDocument


def _labels(language: str) -> tuple[str, str, str, str, str]:
    if language == "pt-BR":
        return ("Resumo profissional", "Competências", "Experiência profissional", "Formação", "Atual")
    return ("Summary", "Skills", "Experience", "Education", "Present")


def render_text(document: ResumeDocument) -> str:
    summary_label, skills_label, experience_label, education_label, present_label = _labels(document.language)
    lines = [document.header.get("name", ""), document.header.get("email", ""), document.header.get("location", ""), "", summary_label, document.summary, "", skills_label, ", ".join(document.skills), "", experience_label]
    for row in document.experience:
        lines.extend([f"{row['role']} | {row['company']} | {row['start_date']} - {row['end_date'] or present_label}"])
        lines.extend(f"- {bullet}" for bullet in row.get("bullets", []))
        lines.append("")
    if document.education:
        lines.append(education_label)
        for row in document.education:
            lines.extend([f"{row['credential']}, {row['field_of_study']} | {row['institution']}", ""])
    return "\n".join(lines).strip() + "\n"


def render_docx(document: ResumeDocument, path: str | Path) -> Path:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - dependency supplied in production
        raise RuntimeError("python-docx is required to render the resume") from exc
    target = Path(path)
    target.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(document.header.get("name", ""), level=0)
    contact = " | ".join(value for value in (document.header.get("email", ""), document.header.get("phone", ""), document.header.get("location", "")) if value)
    if contact:
        doc.add_paragraph(contact)
    summary_label, skills_label, experience_label, education_label, present_label = _labels(document.language)
    doc.add_heading(summary_label, level=1)
    doc.add_paragraph(document.summary)
    doc.add_heading(skills_label, level=1)
    doc.add_paragraph(", ".join(document.skills))
    doc.add_heading(experience_label, level=1)
    for row in document.experience:
        doc.add_heading(f"{row['role']} | {row['company']}", level=2)
        doc.add_paragraph(f"{row['start_date']} - {row['end_date'] or present_label}")
        for bullet in row.get("bullets", []):
            doc.add_paragraph(str(bullet), style="List Bullet")
    if document.education:
        doc.add_heading(education_label, level=1)
        for row in document.education:
            doc.add_paragraph(f"{row['credential']}, {row['field_of_study']} | {row['institution']}")
    doc.save(target)
    return target


def render_pdf_from_docx(docx_path: str | Path, pdf_path: str | Path) -> Path:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice is None:
        raise RuntimeError("libreoffice is required to render the resume PDF")
    source = Path(docx_path)
    target = Path(pdf_path)
    target.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    result = subprocess.run([libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(target.parent), str(source)], capture_output=True, text=True, timeout=60, check=False)
    generated = target.parent / f"{source.stem}.pdf"
    if result.returncode != 0 or not generated.exists():
        raise RuntimeError(f"resume PDF render failed: {result.stderr.strip() or result.stdout.strip()}")
    if generated != target:
        generated.replace(target)
    return target
